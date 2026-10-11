"""结构化简历 → .docx。中文字体必须同时设 w:eastAsia，否则 Word 回落宋体/乱码。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app.agents.schemas import ResumeDraft

CN_FONT = "微软雅黑"


def _set_cn_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = CN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    run.font.size = Pt(size)
    run.font.bold = bold


def _para(doc, text: str, *, size: float = 10.5, bold: bool = False, space_after: float = 2.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_cn_font(run, size=size, bold=bold)
    return p


def _section(doc, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    _set_cn_font(run, size=11, bold=True)


def _join(*parts: str) -> str:
    return " ｜ ".join(x for x in parts if x)


def render_resume_docx(draft: ResumeDraft, out_path: str | Path) -> str:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Pt(48)
        section.bottom_margin = Pt(48)
        section.left_margin = Pt(50)
        section.right_margin = Pt(50)

    _para(doc, draft.name or "（姓名）", size=16, bold=True, space_after=4)
    contact = _join(draft.phone, draft.email)
    if contact:
        _para(doc, contact, size=10)
    if draft.target_role:
        _para(doc, f"求职意向：{draft.target_role}", bold=True)

    if draft.summary:
        _section(doc, "个人总结")
        _para(doc, draft.summary)

    if draft.education:
        _section(doc, "教育背景")
        for ed in draft.education:
            date = _join(f"{ed.start} - {ed.end}") if (ed.start or ed.end) else ""
            _para(doc, _join(date, ed.school, ed.major, ed.degree), bold=True, space_after=1)
            for h in ed.highlights:
                _para(doc, f"· {h}")

    if draft.skills:
        _section(doc, "专业技能")
        for s in draft.skills:
            _para(doc, f"· {s}")

    if draft.internships:
        _section(doc, "实习经历")
        for it in draft.internships:
            date = _join(f"{it.start} - {it.end}") if (it.start or it.end) else ""
            _para(doc, _join(it.org, it.role, date), bold=True, space_after=1)
            for b in it.bullets:
                _para(doc, f"· {b}")

    if draft.projects:
        _section(doc, "项目经历")
        for p in draft.projects:
            _para(doc, _join(p.name, p.role), bold=True, space_after=1)
            for b in p.bullets:
                _para(doc, f"· {b}")

    if draft.awards:
        _section(doc, "获奖情况")
        for a in draft.awards:
            _para(doc, f"· {a}")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return str(out_path)
